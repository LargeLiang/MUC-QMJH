"""Convert verified LaTeX manuscript to editable Word using Pandoc and OOXML.

Run with the Codex bundled Python. Cross references use the final XeLaTeX AUX.
"""
from pathlib import Path
import re, subprocess
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
HERE=ROOT/'基于大语言模型输出文本的选择偏好研究/draft_v3'
DEST=ROOT/'output/doc/数学的实践与认识_审稿修订_v3.docx'
PANDOC=Path('C:/Code/Anaconda3/Library/bin/pandoc.exe')


def main():
    source=(HERE/'manuscript.tex').read_text(encoding='utf-8')
    aux=(HERE/'build/manuscript.aux').read_text(encoding='utf-8')
    labels=dict(re.findall(r'\\newlabel\{([^}]+)\}\{\{([^}]+)\}',aux))
    cites=dict(re.findall(r'\\bibcite\{([^}]+)\}\{([^}]+)\}',aux))
    source=re.sub(r'\\eqref\{([^}]+)\}',lambda m:'('+labels[m[1]]+')',source)
    source=re.sub(r'\\ref\{([^}]+)\}',lambda m:labels[m[1]],source)
    source=re.sub(r'\\cite\{([^}]+)\}',lambda m:'['+','.join(cites[k] for k in m[1].split(','))+']',source)
    source=source.replace('\\OR','\\mathrm{OR}')
    # Materialize caption numbers consistently with the final AUX references.
    for kind in ['table','figure']:
        def caption_number(m):
            block=m[0]
            lab=re.search(r'\\label\{([^}]+)\}',block)
            if lab:
                number=labels[lab[1]]
                block=block.replace('\\caption{','\\caption{'+('表' if kind=='table' else '图')+number+' ',1)
            return block
        source=re.sub(r'\\begin\{'+kind+r'\}.*?\\end\{'+kind+r'\}',caption_number,source,flags=re.S)
    # Explicit equation numbers preserve the manuscript's numbered references.
    parts=re.split(r'(\\section\{[^}]+\})',source);section=0
    for i,part in enumerate(parts):
        if part.startswith('\\section{'):section+=1;continue
        counter=[0]
        def number(m):
            counter[0]+=1
            body=re.sub(r'\\label\{[^}]+\}','',m[1]).strip()
            return '\\['+body+rf'\qquad\text{{({section}.{counter[0]})}}'+'\\]'
        parts[i]=re.sub(r'\\begin\{equation\}(.*?)\\end\{equation\}',number,part,flags=re.S)
    source=''.join(parts)
    source=source.replace('\\begin{tabularx}{\\textwidth}{p{2.8cm}Y Y}','\\begin{tabular}{lll}').replace('\\end{tabularx}','\\end{tabular}')
    source=source.replace('\\FloatBarrier','').replace('\\heiti ','')
    source=re.sub(r'\\includegraphics\[[^]]*\]\{([^}]+)\.pdf\}',r'\\includegraphics[width=15cm]{assets/\1.png}',source)
    # Pandoc cannot use the custom OR command definition after textual expansion.
    source=re.sub(r'\\newcommand\{\\mathrm\{OR\}\}\{[^\n]+\}\n','',source)
    temp=HERE/'build/word_source.tex';temp.write_text(source,encoding='utf-8')
    DEST.parent.mkdir(parents=True,exist_ok=True)
    subprocess.run([str(PANDOC),str(temp),'-f','latex','-t','docx','--standalone','--number-sections',
                    '--resource-path',str(HERE),'-o',str(DEST)],check=True,cwd=HERE)
    doc=Document(DEST)
    for sec in doc.sections:
        sec.page_width=Cm(21);sec.page_height=Cm(29.7)
        sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.2);sec.left_margin=Cm(2.25);sec.right_margin=Cm(2.25)
    for style in doc.styles:
        if style.type not in (1,2):continue
        style.font.name='Times New Roman';style.font.size=Pt(10.5);style.font.color.rgb=RGBColor(0,0,0)
        rp=style.element.get_or_add_rPr();rf=rp.find(qn('w:rFonts'))
        if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
        rf.set(qn('w:eastAsia'),'宋体')
        for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme']:
            rf.attrib.pop(qn('w:'+attr),None)
    for name,size in [('Title',17),('Subtitle',12),('Heading 1',13),('Heading 2',11.5)]:
        st=doc.styles[name];st.font.size=Pt(size);st.font.bold=name!='Subtitle'
        st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6)
    for para in doc.paragraphs:
        fmt=para.paragraph_format;fmt.line_spacing=1.2;fmt.space_after=Pt(4)
        if para.style.name in ('Normal','Body Text','First Paragraph'):
            fmt.first_line_indent=Cm(.74)
        if para.style.name=='Title':
            para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if para.style.name.startswith('Heading'):fmt.keep_with_next=True
    for table in doc.tables:
        table.autofit=False
        n=len(table.columns);widths={3:[2.8,6.8,6.8],4:[3.5,4,4.5,4.4],5:[2.5,3,3,5.4,2.5]}.get(n,[16.4/n]*n)
        for rowi,row in enumerate(table.rows):
            for j,cell in enumerate(row.cells):
                cell.width=Cm(widths[j]);tc=cell._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                tc.append(borders)
                margin=OxmlElement('w:tcMar')
                for edge in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');margin.append(e)
                tc.append(margin)
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.1
                    for run in p.runs:run.font.size=Pt(9)
                if rowi==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEEEEE');tc.append(shade)
            if rowi==0:
                prop=row._tr.get_or_add_trPr();prop.append(OxmlElement('w:tblHeader'))
            # Keep each row together but allow a long table to span pages.
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    # Page numbers in the footer.
    for sec in doc.sections:
        p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    doc.save(DEST)
    print(DEST)


if __name__=='__main__':main()
